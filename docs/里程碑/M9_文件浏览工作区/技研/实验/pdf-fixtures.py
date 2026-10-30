#!/usr/bin/env python3
"""M9-GAP-05：生成 PDF 实验样本（全部合成，固定内容，可重复）。

输出目录里：
  normal.pdf      3 页，中文/英文正文，页内链接、外部链接、书签大纲（LibreOffice 从 DOCX 导出并自带大纲，再用 pypdf 补两个链接）
  many-200.pdf    200 页（normal 前两页循环），用于性能
  encrypted.pdf   用户密码 m9-user（AES-256），用于密码流程
  truncated.pdf   normal.pdf 截掉后半（xref 缺失）
  garbage.pdf     以 %PDF- 开头但内容是随机字节
  not-pdf.pdf     纯文本，扩展名冒充
  bigpage.pdf     单页 14400×14400 pt（PDF 允许的上限量级），测画布内存保护
manifest.json 记录每个文件的 bytes 与 sha256。依赖：LibreOffice、python-docx、pypdf、pycryptodome（均为本机已有，不安装）。
"""
import argparse, hashlib, json, pathlib, shutil, subprocess, sys, tempfile
from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import Link
from pypdf.generic import RectangleObject


def soffice():
    p = shutil.which('soffice') or '/Applications/LibreOffice.app/Contents/MacOS/soffice'
    if not pathlib.Path(p).exists(): raise SystemExit('需要本机已有的 LibreOffice')
    return p


def base_pdf(tmp):
    doc = Document()
    doc.add_heading('文件浏览工作区 Overview', 0)
    doc.add_paragraph('项目维护根目录；会话保存面板与标签状态；文件身份由宿主与路径共同确定。中文与 English mixed text 12345。')
    doc.add_paragraph('PAGE_ONE_MARKER 第一页的结尾，下一页开始是交互。')
    doc.add_page_break()
    doc.add_heading('交互与状态', 1)
    doc.add_paragraph('PAGE_TWO_MARKER 单击预览，双击固定；保存时检查外部变化，冲突保留草稿。')
    doc.add_page_break()
    doc.add_heading('验收与交付', 1)
    doc.add_paragraph('PAGE_THREE_MARKER 全部用例通过才算开发完成。')
    docx = tmp / 'normal.docx'; doc.save(docx)
    out = tmp / 'pdf'; out.mkdir()
    subprocess.run([soffice(), '-env:UserInstallation=' + (tmp / 'profile').as_uri(), '--headless', '--nologo', '--nodefault', '--norestore', '--convert-to', 'pdf', '--outdir', str(out), str(docx)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
    return out / 'normal.pdf'


def with_links_and_outline(src, dest):
    """LibreOffice 已导出大纲，这里只补链接。内部链接用标准的页对象引用 [pageRef /Fit]（pypdf 的 target_page_index 写整数页码，不是真实 PDF 的写法）。"""
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, NumberObject
    r = PdfReader(src); w = PdfWriter(clone_from=r)
    assert len(w.pages) == 3, len(w.pages)
    w.add_annotation(page_number=0, annotation=Link(rect=(72, 600, 300, 620), url='https://example.invalid/external'))
    target = w.pages[2].indirect_reference
    link = DictionaryObject({NameObject('/Type'): NameObject('/Annot'), NameObject('/Subtype'): NameObject('/Link'),
        NameObject('/Rect'): ArrayObject([NumberObject(72), NumberObject(560), NumberObject(300), NumberObject(580)]),
        NameObject('/Border'): ArrayObject([NumberObject(0), NumberObject(0), NumberObject(0)]),
        NameObject('/Dest'): ArrayObject([target, NameObject('/Fit')])})
    page = w.pages[0]; annots = page.get('/Annots'); annots = annots.get_object() if annots is not None else ArrayObject()
    annots.append(w._add_object(link)); page[NameObject('/Annots')] = annots
    with open(dest, 'wb') as f: w.write(f)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--out', required=True); a = ap.parse_args()
    out = pathlib.Path(a.out); out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='m9-pdf-fixtures-') as t:
        tmp = pathlib.Path(t); raw = base_pdf(tmp)
        normal = out / 'normal.pdf'; with_links_and_outline(raw, normal)
        r = PdfReader(normal)
        w = PdfWriter()
        for i in range(200): w.add_page(r.pages[i % 2])
        with open(out / 'many-200.pdf', 'wb') as f: w.write(f)
        w = PdfWriter(clone_from=r); w.encrypt(user_password='m9-user', owner_password='m9-owner', algorithm='AES-256')
        with open(out / 'encrypted.pdf', 'wb') as f: w.write(f)
        data = normal.read_bytes(); (out / 'truncated.pdf').write_bytes(data[:len(data) // 2])
        import random; rnd = random.Random(20261017)
        (out / 'garbage.pdf').write_bytes(b'%PDF-1.7\n' + bytes(rnd.getrandbits(8) for _ in range(4096)))
        (out / 'not-pdf.pdf').write_bytes('这不是 PDF，只是扩展名冒充。\n'.encode('utf-8'))
        w = PdfWriter(); w.add_blank_page(14400, 14400)
        with open(out / 'bigpage.pdf', 'wb') as f: w.write(f)
    manifest = {p.name: {'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(out.glob('*.pdf'))}
    (out / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v['bytes'] for k, v in manifest.items()}, ensure_ascii=False))


if __name__ == '__main__':
    sys.exit(main())
