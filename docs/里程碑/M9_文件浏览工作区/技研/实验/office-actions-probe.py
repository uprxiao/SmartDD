#!/usr/bin/env python3
"""Synthetic OOXML sheet actions; no user files or production editor integration.
Uses tokenized direct references, not string replacement over formulas or ZIP files.
Unsupported dependency forms fail before candidate generation. Untouched parts retain bytes.
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import pathlib
import platform
import posixpath
import re
import shutil
import subprocess
import tempfile
import time
from urllib.parse import unquote
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from lxml import etree as ET

HERE = pathlib.Path(__file__).resolve()
spec = importlib.util.spec_from_file_location('m9_structure', HERE.with_name('office-structure-probe.py'))
base = importlib.util.module_from_spec(spec); spec.loader.exec_module(base)
S, O, R, C = base.S, base.O, base.R, base.C
tree, xml, digest = base.tree, base.xml, base.digest
W = 'xl/workbook.xml'; WR = 'xl/_rels/workbook.xml.rels'
FORMULA_NAMES = {'f', 'formula', 'formula1', 'formula2', 'definedName'}


def sheet_nodes(parts):
    wb = tree(parts[W]); return wb, wb.find('{%s}sheets' % S)


def formula_parts(parts):
    for name, data in parts.items():
        if name == W or name.startswith(('xl/worksheets/', 'xl/charts/')) and name.endswith('.xml'):
            yield name, tree(data)


def reference_token(value, old_name, new_name):
    if '!' not in value:
        return value
    head, cell = value.rsplit('!', 1)
    if '[' in head or ']' in head:
        return value  # Another workbook's name is not this workbook's sheet identity.
    decoded = head[1:-1].replace("''", "'") if head.startswith("'") and head.endswith("'") else head
    if ':' in decoded:
        raise ValueError('unsupported 3D sheet reference: ' + value)
    if decoded.casefold() != old_name.casefold():
        return value
    return '#REF!' if new_name is None else "'%s'!%s" % (new_name.replace("'", "''"), cell)


def rewrite_formula(text, old_name, new_name):
    prefixed = text.startswith('='); tokens = Tokenizer(text if prefixed else '=' + text).items
    values = []
    for token in tokens:
        if token.type == 'FUNC' and token.value.upper().startswith('INDIRECT('):
            raise ValueError('dynamic INDIRECT dependency requires a richer semantic adapter')
        value = token.value
        if token.type == 'OPERAND' and token.subtype == 'RANGE':
            value = reference_token(value, old_name, new_name)
        values.append(value)
    return ('=' if prefixed else '') + ''.join(values)


def preflight(parts):
    if 'xl/calcChain.xml' in parts:
        raise ValueError('calcChain requires a registered invalidation plan')
    for name, doc in formula_parts(parts):
        for node in doc.iter():
            local = ET.QName(node).localname
            if local == 'f' and node.get('t') in ('array', 'shared', 'dataTable'):
                raise ValueError('unsupported formula storage ' + node.get('t'))
            if local in FORMULA_NAMES and node.text:
                rewrite_formula(node.text, '__none__', '__none__')


def rewrite_dependencies(parts, old_name, new_name):
    out = dict(parts); edits = []
    for name, doc in formula_parts(out):
        dirty = False; invalid_formula_nodes = set()
        for node in doc.iter():
            if ET.QName(node).localname not in FORMULA_NAMES or not node.text:
                continue
            previous = node.text; current = rewrite_formula(previous, old_name, new_name)
            if current != previous:
                if new_name is None:
                    # The limited deletion adapter proves direct arithmetic/SUM
                    # error propagation only. IFERROR/IF/other functions need a
                    # full dependency evaluator; do not stamp a false error cache.
                    functions = [t.value.upper() for t in Tokenizer(previous if previous.startswith('=') else '='+previous).items if t.type=='FUNC' and t.subtype=='OPEN']
                    if any(f != 'SUM(' for f in functions):
                        raise ValueError('unsupported deletion error propagation: '+previous)
                    invalid_formula_nodes.add(node)
                node.text = current; dirty = True
                edits.append({'part': name, 'before': previous, 'after': current})
        if dirty:
            # Recompute formula cells in a separate calculation runtime. Never show old v as current.
            for cell in doc.findall('.//{%s}c' % S):
                formula = cell.find('{%s}f' % S)
                if formula is not None:
                    value = cell.find('{%s}v' % S)
                    if value is not None: cell.remove(value)
                    if formula in invalid_formula_nodes:
                        # Matches independently generated LO OOXML after sheet deletion.
                        # This cache follows a proven invalid-reference expression,
                        # not a guessed value for arbitrary formulas.
                        cell.set('t','e'); ET.SubElement(cell,'{%s}v'%S).text='#REF!'
            out[name] = xml(doc)
    wb = tree(out[W]); calc = wb.find('{%s}calcPr' % S)
    if calc is None: calc = ET.SubElement(wb, '{%s}calcPr' % S)
    calc.set('fullCalcOnLoad', '1'); calc.set('forceFullCalc', '1'); out[W] = xml(wb)
    return out, edits


def rename_sheet(parts, old_name, new_name):
    preflight(parts); wb, sheets = sheet_nodes(parts)
    if not new_name or len(new_name.encode('utf-16-le')) // 2 > 31 or re.search(r'[\\/?*\[\]:]', new_name) or new_name.startswith("'") or new_name.endswith("'"):
        raise ValueError('invalid sheet name')
    if any(s.get('name').casefold() == new_name.casefold() and s.get('name') != old_name for s in sheets):
        raise ValueError('duplicate sheet name')
    target = next(s for s in sheets if s.get('name') == old_name)
    target.set('name', new_name); out = dict(parts); out[W] = xml(wb)
    return rewrite_dependencies(out, old_name, new_name)


def reorder_sheets(parts, order):
    preflight(parts); wb, sheets = sheet_nodes(parts); names = [s.get('name') for s in sheets]
    if len(order) != len(names) or set(order) != set(names): raise ValueError('not a full permutation')
    lookup = {s.get('name'): s for s in sheets}; mapping = {i: order.index(name) for i, name in enumerate(names)}
    for node in wb.findall('.//{%s}definedName' % S):
        if node.get('localSheetId') is not None: node.set('localSheetId', str(mapping[int(node.get('localSheetId'))]))
    for view in wb.findall('.//{%s}workbookView' % S):
        view.set('activeTab', str(mapping[int(view.get('activeTab', '0'))]))
        view.set('firstSheet', str(mapping[int(view.get('firstSheet', '0'))]))
    sheets[:] = [lookup[name] for name in order]; out = dict(parts); out[W] = xml(wb)
    return out


def relationship_targets(name, data):
    directory = '' if name == '_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
    for rel in tree(data):
        if rel.get('TargetMode') == 'External': continue
        value = unquote(rel.get('Target', '').split('#')[0])
        yield value.lstrip('/') if value.startswith('/') else posixpath.normpath(posixpath.join(directory, value))


def rels_for(part):
    return posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')


def reachable(parts, start):
    seen = set(); pending = list(start)
    while pending:
        part = pending.pop()
        if part in seen or part not in parts: continue
        seen.add(part)
        rels = part if part.endswith('.rels') else rels_for(part)
        if rels in parts:
            seen.add(rels); pending.extend(relationship_targets(rels, parts[rels]))
    return seen


def delete_sheet(parts, name):
    preflight(parts); wb, sheets = sheet_nodes(parts); names = [s.get('name') for s in sheets]
    index = names.index(name); target = sheets[index]
    if len(sheets) <= 1 or not any(s is not target and s.get('state', 'visible') == 'visible' for s in sheets):
        raise ValueError('must retain a visible worksheet')
    rels = tree(parts[WR]); rel = next(r for r in rels if r.get('Id') == target.get('{%s}id' % O))
    sheet_path = rel.get('Target').lstrip('/') if rel.get('Target').startswith('/') else posixpath.normpath(posixpath.join('xl', rel.get('Target')))
    candidate_deletions = reachable(parts, [sheet_path]); rels.remove(rel); sheets.remove(target)
    for node in list(wb.findall('.//{%s}definedName' % S)):
        local = node.get('localSheetId')
        if local is not None:
            if int(local) == index: node.getparent().remove(node)
            elif int(local) > index: node.set('localSheetId', str(int(local)-1))
    for view in wb.findall('.//{%s}workbookView' % S):
        for key in ['activeTab', 'firstSheet']:
            old = int(view.get(key, '0')); value = old-1 if old > index else min(old, len(sheets)-1)
            view.set(key, str(value))
    out = dict(parts); out[W] = xml(wb); out[WR] = xml(rels)
    still_reachable = reachable(out, ['_rels/.rels'])
    removed = candidate_deletions - still_reachable
    for part in removed: del out[part]
    types = tree(out['[Content_Types].xml'])
    for node in list(types):
        if node.get('PartName', '').lstrip('/') in removed: types.remove(node)
    out['[Content_Types].xml'] = xml(types)
    out, edits = rewrite_dependencies(out, name, None)
    return out, edits, removed


def fixture(root):
    wb = Workbook(); data = wb.active; data.title = 'Data'; calc = wb.create_sheet('Calc'); spare = wb.create_sheet('Spare')
    data.append(['Label', 'Value']); data.append(['A', 3]); data.append(['B', 4]); data['A5'] = 'MERGED'; data.merge_cells('A5:B5')
    data['A2'].comment = Comment('PRESERVE_COMMENT', 'M9 Synthetic')
    chart = BarChart(); chart.add_data(Reference(data, min_col=2, min_row=1, max_row=3), titles_from_data=True); data.add_chart(chart, 'E2')
    wb.defined_names.add(DefinedName('Choices', attr_text="'Data'!$A$2:$A$3"))
    calc.defined_names.add(DefinedName('LocalResult', attr_text="'Calc'!$B$2"))
    spare.defined_names.add(DefinedName('LocalSpare', attr_text="'Spare'!$A$1"))
    calc.append(['RESULT_MARKER', 'Result']); calc['A2'] = 'DIRECT_RESULT'; calc['B2'] = '=Data!$B$2+\'Data\'!B3'
    calc['A3'] = 'SUM_RESULT'; calc['B3'] = '=SUM(Data!B2:B3)'
    calc['A4'] = 'STRING_RESULT'; calc['B4'] = '=IF("Data!B2"="Data!B2",100,0)'
    calc['A5'] = 'LITERAL Data!B2'; calc['A6'] = 'M9_ACTION_FIXTURE'; spare['A1'] = 'SPARE'
    validation = DataValidation(type='list', formula1='Choices'); calc.add_data_validation(validation); validation.add('D2:D3')
    calc.column_dimensions['A'].width = 30; calc.column_dimensions['B'].width = 18; wb.active = 1
    path = root/'actions-source.xlsx'; wb.save(path)
    entries = base.old.add_part(base.old.read_zip(path), base.old.OPAQUE_NAME, base.old.OPAQUE, 'application/octet-stream', ('_rels/.rels', base.old.OPAQUE_NAME), 'urn:smartdd:m9:opaque')
    base.old.write_zip(path, entries); return path


def changed_parts(before, after):
    return {p for p in before.keys() & after.keys() if before[p] != after[p]}


def save_candidate(root, name, before, after, allowed_changed, removed=()):
    actual = changed_parts(before, after)
    assert actual <= set(allowed_changed), ('unexpected changed parts', actual)
    guard = base.checked_mutation(before, after, actual, (), removed)
    out = root/name; base.old.write_zip(out, list(after.items()))
    reader = load_workbook(out, data_only=False)
    return out, reader, guard


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--output', required=True); ap.add_argument('--skip-libreoffice', action='store_true'); args = ap.parse_args()
    report = {'revision':'R3.4', 'startedAt':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()), 'environment':platform.platform(), 'scriptSha256':digest(HERE.read_bytes()), 'scope':'Synthetic sheet action dependency proof, not complete Office engine or production save', 'cases':[]}
    def case(cid, title, fn):
        start=time.monotonic()
        try: facts=fn(); status='passed'
        except Exception as error: facts={'error':repr(error)}; status='failed'
        report['cases'].append({'id':cid,'title':title,'status':status,'durationMs':round((time.monotonic()-start)*1000,2),'facts':facts})
        print(cid,status,title,flush=True)
    with tempfile.TemporaryDirectory(prefix='m9-office-actions-') as temp:
        root=pathlib.Path(temp); original=fixture(root); before=dict(base.old.read_zip(original)); original_hash=digest(original.read_bytes())
        renamed, renamed_edits=rename_sheet(before,'Data',"计划 O'Brien")
        reordered=reorder_sheets(before,['Calc','Spare','Data'])
        deleted, deleted_edits, deleted_parts=delete_sheet(before,'Data')
        allowed={W,'xl/worksheets/sheet2.xml','xl/charts/chart1.xml'}
        paths={}
        def rename_check():
            out,w,g=save_candidate(root,'renamed.xlsx',before,renamed,allowed);paths['rename']=out
            assert w.sheetnames==["计划 O'Brien",'Calc','Spare']
            assert w['Calc']['B2'].value=="='计划 O''Brien'!$B$2+'计划 O''Brien'!B3"
            assert w['Calc']['B4'].value=='=IF("Data!B2"="Data!B2",100,0)'
            assert w['Calc']['A5'].value=='LITERAL Data!B2'
            assert w.defined_names['Choices'].attr_text=="'计划 O''Brien'!$A$2:$A$3"
            assert len(w["计划 O'Brien"]._charts)==1 and w["计划 O'Brien"]['A2'].comment.text=='PRESERVE_COMMENT'
            chart=tree(renamed['xl/charts/chart1.xml']); assert any("计划 O''Brien" in (n.text or '') for n in chart.iter() if ET.QName(n).localname=='f')
            assert renamed[WR]==before[WR] and renamed['[Content_Types].xml']==before['[Content_Types].xml']
            return {**g,'dependencyEdits':renamed_edits,'stringLiteralsPreserved':True,'stableSheetIdsAndRelationships':True}
        case('OA01','sheet改名同步直接公式、命名区域与图表，保留字面量/格式/部件',rename_check)
        def reorder_check():
            out,w,g=save_candidate(root,'reordered.xlsx',before,reordered,{W});paths['reorder']=out
            assert w.sheetnames==['Calc','Spare','Data'] and w.active.title=='Calc'
            assert 'LocalResult' in w['Calc'].defined_names and 'LocalSpare' in w['Spare'].defined_names
            assert w['Calc']['B2'].value==load_workbook(original)['Calc']['B2'].value
            a,b=sheet_nodes(before)[1],sheet_nodes(reordered)[1]
            assert {s.get('name'):(s.get('sheetId'),s.get('{%s}id'%O)) for s in a}=={s.get('name'):(s.get('sheetId'),s.get('{%s}id'%O)) for s in b}
            return {**g,'order':w.sheetnames,'active':w.active.title,'localDefinedNamesRemainOnSameIdentity':True}
        case('OA02','sheet重排保留身份并重映射localSheetId/activeTab',reorder_check)
        def delete_spare():
            out,edits,removed=delete_sheet(before,'Spare');_,w,g=save_candidate(root,'deleted-spare.xlsx',before,out,{W,WR,'[Content_Types].xml'},removed)
            assert w.sheetnames==['Data','Calc'] and removed=={'xl/worksheets/sheet3.xml'}
            assert 'LocalSpare' not in str(out[W]); assert out['xl/charts/chart1.xml']==before['xl/charts/chart1.xml']
            return {**g,'deletedLocalName':True,'dependencyEdits':edits}
        case('OA03','删除无引用工作表只删除其私有部件与本地名称',delete_spare)
        def delete_referenced():
            out,w,g=save_candidate(root,'deleted-data.xlsx',before,deleted,{W,WR,'[Content_Types].xml','xl/worksheets/sheet2.xml'},deleted_parts);paths['delete']=out
            assert w.sheetnames==['Calc','Spare'] and w.active.title=='Calc'
            assert w['Calc']['B2'].value=='=#REF!+#REF!' and w['Calc']['B3'].value=='=SUM(#REF!)'
            assert w.defined_names['Choices'].attr_text=='#REF!'
            assert 'LocalResult' in w['Calc'].defined_names and 'LocalSpare' in w['Spare'].defined_names
            assert 'xl/charts/chart1.xml' in deleted_parts and base.old.OPAQUE_NAME not in deleted_parts
            return {**g,'dependencyEdits':deleted_edits,'invalidReferencesAreExplicit':True}
        case('OA04','删除被直接引用工作表产生明确REF错误并清理专属图表/批注关系',delete_referenced)
        def negative_names():
            rejected=[]
            for name in ['', 'calc', 'bad/name', 'x'*32, "'leading", "trailing'"]:
                try:rename_sheet(before,'Data',name)
                except ValueError:rejected.append(name)
            assert len(rejected)==6;return {'rejected':rejected}
        case('OA05','空/重复/非法/过长名称拒绝，不产生部分改名',negative_names)
        def last_visible():
            wb,sheets=sheet_nodes(before)
            for s in sheets:
                if s.get('name')!='Data':s.set('state','hidden')
            variant=dict(before);variant[W]=xml(wb)
            try:delete_sheet(variant,'Data')
            except ValueError as error:return {'rejected':str(error)}
            raise AssertionError('last visible worksheet deletion accepted')
        case('OA06','最后可见工作表删除被拒绝',last_visible)
        def unsupported():
            rejected=[]
            for kind,text,attrs in [('3D','SUM(Data:Calc!B2)',{}),('dynamic','INDIRECT("Data!B2")',{}),('shared','Data!B2',{'t':'shared','si':'0'}),('array','Data!B2',{'t':'array','ref':'B2:B3'})]:
                variant=dict(before);doc=tree(variant['xl/worksheets/sheet2.xml']);f=doc.find('.//{%s}f'%S);f.text=text;f.attrib.update(attrs);variant['xl/worksheets/sheet2.xml']=xml(doc)
                try:rename_sheet(variant,'Data','Renamed')
                except ValueError as error:rejected.append({'kind':kind,'reason':str(error)})
            assert len(rejected)==4;return {'rejected':rejected,'capabilityStillRequired':'richer dependency adapter, not removed from M9'}
        case('OA07','3D/动态/shared/array依赖拒绝有限补丁，不伪装完整支持',unsupported)
        def external_ref():
            source='[other.xlsx]Data!A1+Data!A1+"Data!A1"'
            output=rewrite_formula(source,'Data','Renamed')
            assert output=='[other.xlsx]Data!A1+\'Renamed\'!A1+"Data!A1"'
            return {'before':source,'after':output,'externalFetch':False,'scope':'formula token unit; not a fully linked external workbook fixture'}
        case('OA08','同名外部工作簿与字符串不被改为本机sheet',external_ref)
        def extra_loss():
            broken=dict(renamed);del broken[base.old.OPAQUE_NAME]
            try:base.checked_mutation(before,broken,changed_parts(before,renamed))
            except ValueError as error:return {'rejected':str(error)}
            raise AssertionError('unknown part deletion accepted')
        case('OA09','sheet操作仍拒绝非预期未知部件损失',extra_loss)
        def shared_part():
            variant=dict(before);rels=tree(variant[WR]);ET.SubElement(rels,'{%s}Relationship'%R,{'Id':'sharedChartProof','Type':'urn:smartdd:m9:shared-dependency','Target':'charts/chart1.xml'});variant[WR]=xml(rels)
            out,_,removed=delete_sheet(variant,'Data');base.validate_relationships(out)
            assert 'xl/charts/chart1.xml' not in removed and out['xl/charts/chart1.xml']!=variant['xl/charts/chart1.xml']
            assert '#REF!' in out['xl/charts/chart1.xml'].decode()
            return {'sharedPartRetained':True,'chartDependencyMarkedRef':True,'boundary':'synthetic extra internal relationship; not a renderer compatibility claim'}
        case('OA10','仍被其它关系引用的部件不误删除，其依赖显式失效',shared_part)
        def error_capture():
            variant=dict(before);doc=tree(variant['xl/worksheets/sheet2.xml']);doc.find('.//{%s}f'%S).text='IFERROR(Data!B2,0)';variant['xl/worksheets/sheet2.xml']=xml(doc)
            try:delete_sheet(variant,'Data')
            except ValueError as error:return {'rejected':str(error),'reason':'do not guess cached error for a formula that may handle REF'}
            raise AssertionError('complex error capture silently replaced by REF cache')
        case('OA13','删除后的错误捕获语义未知时拒绝，不伪造REF缓存',error_capture)
        def independently_calculate(kind):
            source=paths[kind];target=root/('calc-'+kind);target.mkdir();profile=root/('lo-profile-'+kind)
            cmd=[shutil.which('soffice'),'--headless','-env:UserInstallation='+profile.as_uri(),'--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(target),str(source)]
            done=subprocess.run(cmd,capture_output=True,text=True,timeout=25);assert done.returncode==0,done.stderr
            result=target/source.name;assert result.exists(),done.stdout
            w=load_workbook(result,data_only=True);values=[w['Calc'][cell].value for cell in ['B2','B3','B4']]
            assert values==(['#REF!','#REF!',100] if kind=='delete' else [7,7,100]),values
            return {'independentEngine':'LibreOffice','returnCode':done.returncode,'computedValues':values,'convertedFileUsedOnlyForValidation':True,'candidateHash':digest(source.read_bytes()),'calculatedCopyHash':digest(result.read_bytes())}
        if not args.skip_libreoffice:
            case('OA11','独立LibreOffice重算改名后的跨sheet公式而非只读字符串',lambda:independently_calculate('rename'))
            case('OA12','独立LibreOffice重算删除引用为REF错误且无关字符串公式仍100',lambda:independently_calculate('delete'))
            def engine_candidate_extract():
                engine_path=root/'uno-deleted.xlsx';engine_report_path=root/'uno-deleted.json'
                helper=HERE.with_name('office-actions-uno-candidate.py')
                done=subprocess.run(['/Applications/LibreOffice.app/Contents/Resources/python',str(helper),'--input',str(original),'--candidate',str(engine_path),'--output',str(engine_report_path)],capture_output=True,text=True,timeout=30)
                assert done.returncode==0,(done.stdout,done.stderr)
                engine_report=json.loads(engine_report_path.read_text());engine_parts=dict(base.old.read_zip(engine_path))
                assert engine_report['beforeValues']==[7,7,100]
                assert [cell['error'] for cell in engine_report['engineCells']]==[524,524,0]
                # Resolve the engine's renamed package part through its sheet identity.
                ew,es=sheet_nodes(engine_parts);calc=next(s for s in es if s.get('name')=='Calc')
                er=next(r for r in tree(engine_parts[WR]) if r.get('Id')==calc.get('{%s}id'%O))
                target=er.get('Target');engine_sheet=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl',target))
                engine_xml=tree(engine_parts[engine_sheet]);engine_cells={c.get('r'):c for c in engine_xml.findall('.//{%s}c'%S)}
                # Transplant only the two declared formula/result fields. Do not
                # copy engine styles, worksheet layout, rels or whole ZIP output.
                merged=dict(deleted);doc=tree(merged['xl/worksheets/sheet2.xml'])
                for address in ['B2','B3']:
                    cell=next(c for c in doc.findall('.//{%s}c'%S) if c.get('r')==address);incoming=engine_cells[address]
                    assert incoming.get('t')=='e' and incoming.findtext('{%s}v'%S)=='#REF!'
                    for tag in ['f','v']:
                        old=cell.find('{%s}%s'%(S,tag))
                        if old is not None:cell.remove(old)
                        cell.append(copy.deepcopy(incoming.find('{%s}%s'%(S,tag))))
                    cell.set('t',incoming.get('t'))
                merged['xl/worksheets/sheet2.xml']=xml(doc)
                out,w,g=save_candidate(root,'engine-extracted.xlsx',before,merged,{W,WR,'[Content_Types].xml','xl/worksheets/sheet2.xml'},deleted_parts)
                values=load_workbook(out,data_only=True)
                assert [values['Calc'][c].value for c in ['B2','B3']]==['#REF!','#REF!']
                assert [w['Calc'][c].value for c in ['B2','B3']]==['=#REF!+#REF!','=SUM(#REF!)']
                assert merged[base.old.OPAQUE_NAME]==before[base.old.OPAQUE_NAME]
                hashes={name:digest(data) for name,data in before.items() if name in merged and data==merged[name]}
                engine_loss=sorted(set(before)-set(engine_parts))
                return {**g,'independentReader':'openpyxl formula and cached-error modes; not a recalculation engine','engineReport':engine_report,'engineWholePackagePartsMissing':engine_loss,'engineWholePackageNeverInstalled':True,'retainedOriginalPartHashes':hashes,'candidateOnly':True,'releaseGate':'Referenced-sheet delete remains blocked: OA12 plus independent engine self-roundtrip fail. Typed cache read is not calculation compatibility.'}
            case('OA14','UNO生成候选后仅提取声明公式字段，原ZIP其余部件哈希保全',engine_candidate_extract)
        assert digest(original.read_bytes())==original_hash
        report['originalInputUnchanged']=True
    report['summary']={'passed':sum(c['status']=='passed' for c in report['cases']),'failed':sum(c['status']=='failed' for c in report['cases'])}
    groups={'structuralCandidates':['OA01','OA02','OA03','OA04','OA10'],'rejectionGuards':['OA05','OA06','OA07','OA09','OA13'],'formulaTokenUnit':['OA08'],'independentCalculation':['OA11','OA12'],'engineCandidateExtraction':['OA14']}
    report['evidenceGroups']={name:{'passed':sum(c['id'] in ids and c['status']=='passed' for c in report['cases']),'failed':sum(c['id'] in ids and c['status']=='failed' for c in report['cases']),'cases':ids} for name,ids in groups.items()}
    report['completedAt']=time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime());report['cleanup']='Owned temporary XLSX files, private LibreOffice profiles and converted copies removed'
    report['limitations']=['Direct references only; shared/array/3D/dynamic dependencies require further adapter work','No production UI/save service or Microsoft Office layout acceptance','No general row/column insertion, chart editing, workbook copy, signed VBA or arbitrary OOXML claim','Existing formula calcChain intentionally blocked until invalidation is registered; formula caches rebuilt only on disposable LibreOffice copies']
    pathlib.Path(args.output).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report['summary']));raise SystemExit(report['summary']['failed']>0)


if __name__=='__main__':main()
