#!/usr/bin/env python3
"""R3.4 Office structural editing feasibility, synthetic packages only.

No user documents, Office UI or macros. Preserve untouched ZIP part bytes;
validate exact mutation lists, internal relationships and independent readers.
This is a narrow adapter experiment, not a complete Office editor.
"""
import argparse
import copy
import hashlib
import importlib.util
import json
import pathlib
import platform
import posixpath
import re
import shutil
import subprocess
import tempfile
import time
from urllib.parse import unquote
from lxml import etree as ET

HERE = pathlib.Path(__file__).resolve()
spec = importlib.util.spec_from_file_location('m9_previous_patch', HERE.with_name('ooxml-patch-probe.py'))
old = importlib.util.module_from_spec(spec)
spec.loader.exec_module(old)
R = 'http://schemas.openxmlformats.org/package/2006/relationships'
O = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
C = 'http://schemas.openxmlformats.org/package/2006/content-types'
parser = ET.XMLParser(resolve_entities=False, no_network=True)


def tree(data):
    return ET.fromstring(data, parser)


def xml(node):
    return ET.tostring(node, encoding='UTF-8', xml_declaration=True, standalone=True)


def digest(data):
    return hashlib.sha256(data).hexdigest()


def validate_relationships(parts):
    checked = 0
    for name, data in parts.items():
        if not name.endswith('.rels'):
            continue
        node = tree(data)
        ids = [r.get('Id') for r in node]
        if len(ids) != len(set(ids)):
            raise ValueError('duplicate relationship IDs in ' + name)
        directory = '' if name == '_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
        for rel in node:
            if rel.get('TargetMode') == 'External':
                continue  # Existing links are never fetched.
            target = unquote(rel.get('Target', '').split('#')[0])
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(directory, target))
            if target not in parts:
                raise ValueError('broken relationship %s -> %s' % (name, target))
            checked += 1
    for override in tree(parts['[Content_Types].xml']):
        if override.tag == '{%s}Override' % C and override.get('PartName', '').lstrip('/') not in parts:
            raise ValueError('content type references a missing part')
    return checked


def checked_mutation(before, after, changed=(), added=(), removed=()):
    changed = set(changed); added = set(added); removed = set(removed)
    actual_changed = {k for k in before.keys() & after.keys() if before[k] != after[k]}
    if actual_changed != changed or after.keys() - before.keys() != added or before.keys() - after.keys() != removed:
        raise ValueError('unexpected package mutation: changed=%r added=%r removed=%r' %
                         (actual_changed, after.keys() - before.keys(), before.keys() - after.keys()))
    for name in changed | added:
        if name.endswith('.rels'):
            prior = {ET.tostring(r) for r in tree(before.get(name, ('<Relationships xmlns="%s"/>' % R).encode())) if r.get('TargetMode') == 'External'}
            current = {ET.tostring(r) for r in tree(after[name]) if r.get('TargetMode') == 'External'}
            if current - prior:
                raise ValueError('operation must not introduce an external relationship')
    count = validate_relationships(after)
    return {'changed': sorted(changed), 'added': sorted(added), 'removed': sorted(removed),
            'unchangedPartCount': len(before.keys() & after.keys() - changed), 'internalRelationships': count,
            'opaquePreserved': before[old.OPAQUE_NAME] == after[old.OPAQUE_NAME]}


def append_sheet(before):
    out = dict(before); wb = tree(out['xl/workbook.xml']); sheets = wb.find('{%s}sheets' % S)
    rels = tree(out['xl/_rels/workbook.xml.rels']); types = tree(out['[Content_Types].xml'])
    number = max(int(re.search(r'sheet(\d+)\.xml$', n)[1]) for n in out if re.match(r'xl/worksheets/sheet\d+\.xml$', n)) + 1
    target = 'xl/worksheets/sheet%d.xml' % number
    rid = 'm9AddedSheet'
    if any(r.get('Id') == rid for r in rels) or any(s.get('name') == 'Added' for s in sheets):
        raise ValueError('collision')
    sheet_id = max(int(s.get('sheetId')) for s in sheets) + 1
    ET.SubElement(sheets, '{%s}sheet' % S, {'name': 'Added', 'sheetId': str(sheet_id), '{%s}id' % O: rid})
    ET.SubElement(rels, '{%s}Relationship' % R, {'Id': rid, 'Type': O + '/worksheet', 'Target': 'worksheets/sheet%d.xml' % number})
    ET.SubElement(types, '{%s}Override' % C, {'PartName': '/' + target, 'ContentType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
    out['xl/workbook.xml'] = xml(wb); out['xl/_rels/workbook.xml.rels'] = xml(rels); out['[Content_Types].xml'] = xml(types)
    out[target] = ('<worksheet xmlns="%s"><dimension ref="A1:B2"/><cols><col min="1" max="1" width="28" customWidth="1"/></cols><sheetData>'
                   '<row r="1"><c r="A1" t="inlineStr"><is><t>STRUCTURE_ADDED</t></is></c><c r="B1"><v>12</v></c></row>'
                   '<row r="2"><c r="B2"><f>B1*2</f><v>24</v></c></row>'
                   '</sheetData></worksheet>' % S).encode()
    return out, {'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml'}, {target}


def duplicate_slide(before):
    out = dict(before); presentation = tree(out['ppt/presentation.xml']); ids = presentation.find('{%s}sldIdLst' % P)
    rels = tree(out['ppt/_rels/presentation.xml.rels']); types = tree(out['[Content_Types].xml'])
    number = max(int(re.search(r'slide(\d+)\.xml$', n)[1]) for n in out if re.match(r'ppt/slides/slide\d+\.xml$', n)) + 1
    new_slide = 'ppt/slides/slide%d.xml' % number
    new_rels = 'ppt/slides/_rels/slide%d.xml.rels' % number
    out[new_slide] = out['ppt/slides/slide1.xml']
    slide_rels = tree(out['ppt/slides/_rels/slide1.xml.rels'])
    added = {new_slide, new_rels}
    for rel in slide_rels:
        if rel.get('Type') != O + '/notesSlide':
            continue
        source = posixpath.normpath(posixpath.join('ppt/slides', rel.get('Target')))
        source_rels = posixpath.join(posixpath.dirname(source), '_rels', posixpath.basename(source) + '.rels')
        note_n = max(int(re.search(r'notesSlide(\d+)\.xml$', n)[1]) for n in out if re.match(r'ppt/notesSlides/notesSlide\d+\.xml$', n)) + 1
        note = 'ppt/notesSlides/notesSlide%d.xml' % note_n
        note_rels = 'ppt/notesSlides/_rels/notesSlide%d.xml.rels' % note_n
        out[note] = out[source]
        nr = tree(out[source_rels])
        for back in nr:
            if back.get('Type') == O + '/slide':
                back.set('Target', '../slides/slide%d.xml' % number)
        out[note_rels] = xml(nr); added.update([note, note_rels])
        rel.set('Target', '../notesSlides/notesSlide%d.xml' % note_n)
        ET.SubElement(types, '{%s}Override' % C, {'PartName': '/' + note, 'ContentType': 'application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml'})
    out[new_rels] = xml(slide_rels)
    ET.SubElement(ids, '{%s}sldId' % P, {'id': str(max(int(i.get('id')) for i in ids) + 1), '{%s}id' % O: 'm9DuplicatedSlide'})
    ET.SubElement(rels, '{%s}Relationship' % R, {'Id': 'm9DuplicatedSlide', 'Type': O + '/slide', 'Target': 'slides/slide%d.xml' % number})
    ET.SubElement(types, '{%s}Override' % C, {'PartName': '/' + new_slide, 'ContentType': 'application/vnd.openxmlformats-officedocument.presentationml.slide+xml'})
    out['ppt/presentation.xml'] = xml(presentation); out['ppt/_rels/presentation.xml.rels'] = xml(rels); out['[Content_Types].xml'] = xml(types)
    return out, {'ppt/presentation.xml', 'ppt/_rels/presentation.xml.rels', '[Content_Types].xml'}, added


def xlsx_fixture(root):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws = wb.active; ws.title = 'Data'
    ws.append(['Label', 'Value']); ws.append(['A', 3]); ws.append(['B', 4]); ws['C2'] = '=SUM(B2:B3)'
    ws['A5'] = 'MERGED'; ws.merge_cells('A5:B5'); ws['A2'].comment = Comment('PRESERVE_COMMENT', 'Synthetic')
    validation = DataValidation(type='whole', operator='between', formula1=0, formula2=100)
    ws.add_data_validation(validation); validation.add('B2:B3')
    chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True); ws.add_chart(chart, 'E2')
    path = root / 'chart-source.xlsx'; wb.save(path)
    entries = old.add_part(old.read_zip(path), old.OPAQUE_NAME, old.OPAQUE, 'application/octet-stream', ('_rels/.rels', old.OPAQUE_NAME), 'urn:smartdd:m9:opaque')
    old.write_zip(path, entries)
    return path


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--output', required=True); args = ap.parse_args()
    report = {'revision': 'R3.4', 'startedAt': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
              'scriptSha256': digest(HERE.read_bytes()), 'environment': platform.platform(),
              'scope': 'Synthetic OOXML structural mutation prototypes, not production codecs or Microsoft Office acceptance', 'cases': []}
    def case(cid, title, fn):
        started = time.monotonic()
        try: status, facts = 'passed', fn()
        except Exception as error: status, facts = 'failed', {'error': repr(error)}
        report['cases'].append({'id': cid, 'title': title, 'status': status, 'durationMs': round((time.monotonic()-started)*1000, 2), 'facts': facts})
    with tempfile.TemporaryDirectory(prefix='m9-structure-') as temp:
        root = pathlib.Path(temp); xp = xlsx_fixture(root); pp = old.make_pptx(root)
        xb = dict(old.read_zip(xp)); pb = dict(old.read_zip(pp)); original_hashes = {str(p): digest(p.read_bytes()) for p in [xp, pp]}
        xa, xc, xn = append_sheet(xb); pa, pc, pn = duplicate_slide(pb)
        xout = root/'edited.xlsx'; pout = root/'edited.pptx'
        old.write_zip(xout, list(xa.items())); old.write_zip(pout, list(pa.items()))
        def sheet_check():
            from openpyxl import load_workbook
            guard = checked_mutation(xb, xa, xc, xn)
            w = load_workbook(xout); a = w['Added']; data = w['Data']
            assert a['B1'].value == 12 and a['B2'].value == '=B1*2'
            assert data['C2'].value == '=SUM(B2:B3)' and len(data._charts) == 1
            assert 'A5:B5' in data.merged_cells and data['A2'].comment.text == 'PRESERVE_COMMENT'
            assert len(data.data_validations.dataValidation) == 1
            return {**guard, 'independentReader': 'openpyxl', 'sheets': w.sheetnames, 'chartFormulaMergeCommentValidationPreserved': True}
        case('OS01', '新增工作表保留旧表公式、图表、合并、批注、校验与未知部件', sheet_check)
        def slide_check():
            guard = checked_mutation(pb, pa, pc, pn)
            ids = tree(pa['ppt/presentation.xml']).find('{%s}sldIdLst' % P)
            assert len(ids) == 3 and len({n.get('id') for n in ids}) == 3
            assert pa['ppt/slides/slide3.xml'] == pb['ppt/slides/slide1.xml']
            slide_rels = tree(pa['ppt/slides/_rels/slide3.xml.rels'])
            note_rel = next(r for r in slide_rels if r.get('Type') == O + '/notesSlide')
            note_path = posixpath.normpath(posixpath.join('ppt/slides', note_rel.get('Target')))
            assert note_path not in pb and pa[note_path] == pb['ppt/notesSlides/notesSlide1.xml']
            back_path = posixpath.join(posixpath.dirname(note_path), '_rels', posixpath.basename(note_path) + '.rels')
            back_rel = next(r for r in tree(pa[back_path]) if r.get('Type') == O + '/slide')
            assert back_rel.get('Target') == '../slides/slide3.xml'
            return {**guard, 'slideCount': len(ids), 'notesClonedWithNewBackReference': True}
        case('OS02', '复制幻灯片与独立备注，原幻灯片/主题/未知部件不变', slide_check)
        def reorder():
            out = dict(pa); t = tree(out['ppt/presentation.xml']); ids = t.find('{%s}sldIdLst' % P); ids.insert(0, ids[-1]); out['ppt/presentation.xml'] = xml(t)
            facts = checked_mutation(pa, out, {'ppt/presentation.xml'})
            assert ids[0].get('{%s}id' % O) == 'm9DuplicatedSlide'
            return facts
        case('OS03', '幻灯片重排只改presentation顺序，不重写对象和备注', reorder)
        def hidden():
            out = dict(pa); t = tree(out['ppt/slides/slide3.xml']); t.set('show', '0'); out['ppt/slides/slide3.xml'] = xml(t)
            facts = checked_mutation(pa, out, {'ppt/slides/slide3.xml'})
            assert tree(out['ppt/slides/slide3.xml']).get('show') == '0'
            return facts
        case('OS04', '隐藏幻灯片只修改show标记，保留对象内容', hidden)
        def refused(kind):
            out = dict(xa)
            if kind == 'loss': del out[old.OPAQUE_NAME]
            elif kind == 'chart': out['xl/charts/chart1.xml'] += b' '
            elif kind == 'relation':
                t = tree(out['xl/_rels/workbook.xml.rels']); t[-1].set('Target', 'worksheets/missing.xml'); out['xl/_rels/workbook.xml.rels'] = xml(t)
            else:
                t = tree(out['xl/_rels/workbook.xml.rels']); ET.SubElement(t, '{%s}Relationship' % R, {'Id':'injected', 'Type':O+'/externalLink', 'Target':'https://example.invalid/never-requested', 'TargetMode':'External'}); out['xl/_rels/workbook.xml.rels'] = xml(t)
            try: checked_mutation(xb, out, xc, xn)
            except ValueError as error: return {'rejected': True, 'reason': str(error), 'originalUnchanged': digest(xp.read_bytes()) == original_hashes[str(xp)]}
            raise AssertionError('unsafe mutation accepted')
        for cid, kind in [('OS05','loss'),('OS06','chart'),('OS07','relation'),('OS08','external')]:
            case(cid, '拒绝非预期损失/修改/断链/新增外部关系：'+kind, lambda kind=kind: refused(kind))
        def render(path, expected):
            r = old.soffice_pdf(shutil.which('soffice'), path, root)
            assert r['pdfExists'] and r['convert']['exitCode'] == 0, r
            assert expected in r['text'], r
            return {'renderer': 'LibreOffice headless', **r, 'layoutPixelParity': 'not-tested'}
        case('OS09', '独立LibreOffice读取新增工作表并导出PDF', lambda: render(xout,'STRUCTURE_ADDED'))
        case('OS10', '独立LibreOffice读取复制幻灯片并导出PDF', lambda: render(pout,'PPT_TITLE_ONE'))
        assert all(digest(pathlib.Path(p).read_bytes()) == h for p,h in original_hashes.items())
        report['originalInputsUnchanged'] = True
    report['completedAt'] = time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
    report['cleanup'] = 'Temporary packages, private LibreOffice profiles and PDFs removed'
    report['limitations'] = ['No Microsoft Office installation or real user-created Office files used',
        'Only sheet append, slide clone/reorder/hide were implemented; arbitrary row insertion, charts/formula editing and all Office commands remain unproven',
        'Untouched ZIP part bytes are preserved; ZIP container metadata and visual pixel parity are not asserted',
        'Macro-like placeholder bytes are inert, not a signed VBA project compatibility test',
        'Not connected to the WillYou UI, real disk save service or packaged application']
    pathlib.Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'cases': [(c['id'],c['status']) for c in report['cases']], 'output':args.output}, ensure_ascii=False))
    raise SystemExit(any(c['status']=='failed' for c in report['cases']))


if __name__ == '__main__':
    main()
