#!/usr/bin/env python3
"""Use already-installed tools only, with owned synthetic files and private LO profiles.
No real project documents, macros, uploads, package installation or GUI automation.
"""
import argparse, base64, hashlib, importlib.metadata, json, os, pathlib, re
import shutil, struct, subprocess, sys, tempfile, time, zipfile, zlib
import xml.etree.ElementTree as ET

HERE = pathlib.Path(__file__).resolve()
OPAQUE = b'M9 owned preservation sentinel v1\x00\xff\x01'
OPAQUE_NAME = 'm9/opaque-preserve.bin'
NS_REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS_CT = 'http://schemas.openxmlformats.org/package/2006/content-types'

def sha(data): return hashlib.sha256(data).hexdigest()
def run(args, timeout=45):
    started = time.monotonic()
    try:
        p = subprocess.run([str(x) for x in args], capture_output=True, timeout=timeout, check=False)
        return {'exitCode': p.returncode, 'durationMs': round((time.monotonic()-started)*1000, 2),
                'stdout': p.stdout.decode('utf-8','replace').strip(), 'stderr': p.stderr.decode('utf-8','replace').strip()}
    except subprocess.TimeoutExpired:
        return {'exitCode': None, 'timeoutSeconds': timeout, 'durationMs': round((time.monotonic()-started)*1000, 2)}

def package_edit(source, dest, marker=None, inject=False):
    with zipfile.ZipFile(source) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    if marker:
        old, new = marker
        for name, data in list(entries.items()):
            if name.endswith('.xml'): entries[name] = data.replace(old.encode(), new.encode())
    if inject:
        entries[OPAQUE_NAME] = OPAQUE
        # Preserve original OPC namespace spelling; unrelated XML rewriting is not the test.
        entries['_rels/.rels'] = entries['_rels/.rels'].replace(b'</Relationships>',
            ('<Relationship Id="m9Opaque" Type="urn:smartdd:m9:opaque-preservation-probe" Target="'+OPAQUE_NAME+'"/></Relationships>').encode())
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(b'</Types>',
            ('<Override PartName="/'+OPAQUE_NAME+'" ContentType="application/octet-stream"/></Types>').encode())
    with zipfile.ZipFile(dest, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items(): z.writestr(name, data)

def inventory(file):
    with zipfile.ZipFile(file) as z:
        bad = z.testzip()
        assert bad is None, bad
        entries = {n: z.read(n) for n in z.namelist()}
        xml_errors = []
        for name, data in entries.items():
            if name.endswith(('.xml','.rels')):
                try: ET.fromstring(data)
                except ET.ParseError as e: xml_errors.append({'name': name, 'error': str(e)})
        text = ''.join(''.join(ET.fromstring(data).itertext()) for name,data in entries.items() if name.endswith('.xml'))
        return {'bytes': file.stat().st_size, 'sha256': sha(file.read_bytes()), 'entries': sorted(entries),
                'opaquePreserved': entries.get(OPAQUE_NAME) == OPAQUE, 'xmlErrors': xml_errors,
                'editMarkers': sorted(set(re.findall(r'M9_(?:BEFORE|AFTER)_EDIT_(?:DOCX|XLSX|PPTX)', text))),
                'entryHashes': {n:sha(v) for n,v in entries.items()}}

def png_bytes():
    def chunk(t, data): return struct.pack('>I',len(data))+t+data+struct.pack('>I',zlib.crc32(t+data)&0xffffffff)
    raw=b''.join(b'\x00'+bytes([25,120,110,255])*16 for _ in range(16))
    return b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>2I5B',16,16,8,6,0,0,0))+chunk(b'IDAT',zlib.compress(raw))+chunk(b'IEND',b'')

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--output',required=True); args=parser.parse_args()
    report={'schemaVersion':1,'startedAt':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),
            'scriptSha256':sha(HERE.read_bytes()),'scope':'Existing local engines, synthetic OOXML; no UI or production integration',
            'tools':{},'cases':[]}
    for tool in ['soffice','pandoc','pdfinfo','pdftotext']:
        resolved=shutil.which(tool); report['tools'][tool]={'path':resolved}
        if resolved: report['tools'][tool]['version']=run([resolved,'--version' if tool in ['soffice','pandoc'] else '-v'])
    for module in ['openpyxl','python-docx','pypdf']:
        try: report['tools'][module]={'version':importlib.metadata.version(module),'installedBeforeProbe':True}
        except importlib.metadata.PackageNotFoundError: report['tools'][module]={'available':False}
    if not report['tools']['soffice']['path']: raise RuntimeError('Existing LibreOffice required; this script does not install it')
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.comments import Comment
    from openpyxl.chart import BarChart, Reference
    from docx import Document
    from docx.shared import Inches
    from pypdf import PdfReader
    with tempfile.TemporaryDirectory(prefix='m9-office-research-') as temp:
        root=pathlib.Path(temp); fixtures=root/'fixtures'; fixtures.mkdir(); staging=root/'staging'; staging.mkdir()
        image=root/'owned.png'; image.write_bytes(png_bytes())
        doc=Document(); doc.add_heading('M9_BEFORE_EDIT_DOCX',0); doc.add_paragraph('原创合成文档：中文、表格、图片与分页。')
        table=doc.add_table(rows=2,cols=2); table.cell(0,0).text='Feature'; table.cell(0,1).text='Status'; table.cell(1,0).text='保留'; table.cell(1,1).text='待验证'
        doc.add_picture(str(image),width=Inches(.3)); doc.add_page_break(); doc.add_heading('SECOND_PAGE_MARKER',1); doc.save(staging/'sample.docx')
        book=Workbook(); ws=book.active; ws.title='Data'; ws.append(['M9_BEFORE_EDIT_XLSX','Value','Formula','Precision'])
        ws.column_dimensions['A'].width=40; ws.column_dimensions['D'].width=28; ws.append([2,3,'=SUM(A2:B2)','000012345678901234567890']); ws['A1'].font=Font(bold=True,color='FFFFFF'); ws['A1'].fill=PatternFill('solid',fgColor='19786E')
        ws['B2'].number_format='0.00'; ws['A4']='MERGED_MARKER'; ws.merge_cells('A4:B4'); ws['A2'].comment=Comment('M9_COMMENT_MARKER','Synthetic')
        chart=BarChart(); chart.add_data(Reference(ws,min_col=1,max_col=2,min_row=2,max_row=2)); ws.add_chart(chart,'F2')
        second=book.create_sheet('Second'); second['A1']='SECOND_SHEET_MARKER'; book.save(staging/'sample.xlsx')
        markdown=root/'slides.md'; markdown.write_text('# M9_BEFORE_EDIT_PPTX\n\n- 原创合成演示稿\n- 中文与图片\n\n![Owned image](owned.png)\n\n# SECOND_SLIDE_MARKER\n\nNo user documents were read.\n',encoding='utf8')
        p=run([report['tools']['pandoc']['path'],markdown,'--slide-level=1','--resource-path',root,'-t','pptx','-o',staging/'sample.pptx'])
        if p['exitCode']!=0: raise RuntimeError(f'Fixture generation failed: {p}')
        formats={'docx':'docx:Office Open XML Text','xlsx':'xlsx:Calc MS Excel 2007 XML','pptx':'pptx:Impress MS PowerPoint 2007 XML'}
        for extension,format_name in formats.items():
            source=fixtures/f'sample.{extension}'; package_edit(staging/f'sample.{extension}',source,inject=True)
            edited=root/f'edited.{extension}'; package_edit(source,edited,marker=(f'M9_BEFORE_EDIT_{extension.upper()}',f'M9_AFTER_EDIT_{extension.upper()}'))
            before=inventory(edited); record={'format':extension,'fixtureOriginal':inventory(source),'inputEdited':before,'editMechanism':'single owned marker replacement in ZIP XML; all other parts copied intact'}
            output=root/f'roundtrip-{extension}'; output.mkdir(); profile=root/f'profile-{extension}'; pdfout=root/f'pdf-{extension}'; pdfout.mkdir()
            base=[report['tools']['soffice']['path'],f'-env:UserInstallation={profile.as_uri()}','--headless','--nologo','--nodefault','--norestore']
            command=base+['--convert-to',format_name,'--outdir',output,edited]
            record['conversion']=run(command); converted=output/edited.name
            if converted.exists():
                after=inventory(converted); record['outputRoundtrip']=after
                record['missingParts']=sorted(set(before['entries'])-set(after['entries'])); record['newParts']=sorted(set(after['entries'])-set(before['entries']))
                record['unknownPartDropped']=before['opaquePreserved'] and not after['opaquePreserved']
                record['markerPreserved']=f'M9_AFTER_EDIT_{extension.upper()}' in after['editMarkers']
                if extension=='xlsx':
                    parsed=load_workbook(converted,data_only=False); values=load_workbook(converted,data_only=True)
                    record['independentRead']={'reader':'openpyxl','sheets':parsed.sheetnames,'formula':parsed['Data']['C2'].value,'calculatedResult':values['Data']['C2'].value,'precisionText':parsed['Data']['D2'].value,'numberFormat':parsed['Data']['B2'].number_format,'mergeRanges':list(map(str,parsed['Data'].merged_cells.ranges)),'commentText':parsed['Data']['A2'].comment.text if parsed['Data']['A2'].comment else None,'chartParts':len([x for x in after['entries'] if re.match(r'xl/charts/chart\d+\.xml$',x)])}
                elif extension=='docx':
                    parsed=Document(converted); record['independentRead']={'reader':'python-docx','paragraphs':[x.text for x in parsed.paragraphs],'tableCount':len(parsed.tables),'inlinePictureCount':len(parsed.inline_shapes)}
                else:
                    record['independentRead']={'reader':'stdlib ZIP/XML (not a presentation renderer)','slides':len([x for x in after['entries'] if re.match(r'ppt/slides/slide\d+\.xml$',x)]),'imageParts':len([x for x in after['entries'] if x.startswith('ppt/media/')])}
                record['pdfConversion']=run(base+['--convert-to','pdf','--outdir',pdfout,converted]); pdf=pdfout/'edited.pdf'
                if pdf.exists():
                    reader=PdfReader(pdf); text='\n'.join(page.extract_text() or '' for page in reader.pages)
                    record['pdf']={'bytes':pdf.stat().st_size,'sha256':sha(pdf.read_bytes()),'pages':len(reader.pages),'editedTextFound':f'M9_AFTER_EDIT_{extension.upper()}' in text,'chineseTextFound':'中文' in text,'visualLayoutVerified':False}
                    if report['tools']['pdftotext']['path']:
                        independent=run([report['tools']['pdftotext']['path'],pdf,'-']); record['pdf']['popplerTextMarkerFound']=f'M9_AFTER_EDIT_{extension.upper()}' in independent.get('stdout','')
                record['status']='converted-with-preservation-gap' if record['unknownPartDropped'] else 'converted-limited-fixture'
            else: record['status']='conversion-failed'
            # Replace temporary absolute paths in logs; no source document bytes are logged.
            record=json.loads(json.dumps(record).replace(str(root),'<fixture>'))
            report['cases'].append(record)
        report['fixtureCleanup']='completed on exit from TemporaryDirectory'
    report['completedAt']=time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())
    pathlib.Path(args.output).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps({'tools':report['tools'],'cases':[{k:v for k,v in row.items() if k in ['format','status','markerPreserved','unknownPartDropped','missingParts','independentRead','pdf','conversion','pdfConversion']} for row in report['cases']]},ensure_ascii=False,indent=2))
    if any(x['status']=='conversion-failed' for x in report['cases']): sys.exit(1)

if __name__=='__main__': main()
