#!/usr/bin/env python3
"""Pre-development Office proof: generated files only, declared OOXML changes.
Usage: python3 .../predev-office-probe.py --prepare OUTDIR
Then invoke the independent reader/runtime listed in the report. No product saves.
"""
import argparse, copy, hashlib, importlib.util, json, pathlib, platform, posixpath, subprocess, time, zipfile
from io import BytesIO
from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.workbook.defined_name import DefinedName
HERE=pathlib.Path(__file__).resolve()
spec=importlib.util.spec_from_file_location('old_actions',HERE.with_name('office-actions-probe.py'));old=importlib.util.module_from_spec(spec);spec.loader.exec_module(old)
S,O,R,C=old.S,old.O,old.R,old.C
P='http://schemas.openxmlformats.org/presentationml/2006/main';A='http://schemas.openxmlformats.org/drawingml/2006/main';CH='http://schemas.openxmlformats.org/drawingml/2006/chart'
tree,xml=old.tree,old.xml
sha=lambda b:hashlib.sha256(b).hexdigest()
def read(p):return dict(old.base.old.read_zip(p))
def write(p,d):old.base.old.write_zip(p,list(d.items()))
def guard(before,after,changed,added=(),removed=()):
 facts=old.base.checked_mutation(before,after,changed,added,removed)
 facts['retainedHashes']={p:sha(v) for p,v in before.items() if p in after and after[p]==v}
 return facts

def rich_xlsx(root):
 w=Workbook();d=w.active;d.title='Data';c=w.create_sheet('Calc');s=w.create_sheet('Spare')
 d.append(['Label','Value']);d.append(['A',3]);d.append(['B',4]);s['A1']=5
 t=Table(displayName='Sales',ref='A1:B3');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium9',showRowStripes=True);d.add_table(t)
 ch=BarChart();ch.add_data(Reference(d,min_col=2,min_row=1,max_row=3),titles_from_data=True);d.add_chart(ch,'E2')
 w.defined_names.add(DefinedName('InputTotal',attr_text="'Data'!$B$2:$B$3"))
 forms={
 'B2':'=Data!$B$2+Data!B3','B3':'=SUM(Data!B2:B3)','B4':'=IF("Data!B2"="Data!B2",100,0)',
 'B5':'=IFERROR(Data!B2,77)','B6':'=IF(FALSE,Data!B2,42)','B7':'=IFERROR(B2,88)',
 'B8':'=SUM(InputTotal)','B9':'=SUM(Sales[Value])','B10':'=SUM(Data:Spare!A1)',
 'B11':'=Data!B2*2','B12':'=Data!B3*2','B15':'=IFERROR(SUM(Data!B2:B3),99)',
 'B16':'=IFERROR(INDIRECT("Data!B2"),66)'}
 for addr,f in forms.items():c[addr]=f;c['A'+addr[1:]]=addr
 c['D2']=ArrayFormula('D2:D3','=Data!B2:B3*3');c['D3']=0
 c['E2']=123.456;c['E3']=789;c['E2'].font=Font(name='Arial',bold=True,color='FF123456');c['E2'].number_format='0.00';c['E3'].font=copy.copy(c['E2'].font)
 c['G2']='MERGE_TOP';c['H2']='MERGE_COVERED'
 path=root/'rich-source.xlsx';w.save(path);parts=read(path)
 # Genuine shared-formula master/follower encoding; preserving original references.
 doc=tree(parts['xl/worksheets/sheet2.xml']);cells={x.get('r'):x for x in doc.findall('.//{%s}c'%S)}
 for addr in ['B11','B12']:
  f=cells[addr].find('{%s}f'%S);f.set('t','shared');f.set('si','7')
  if addr=='B11':f.set('ref','B11:B12')
  else:f.text=None
 parts['xl/worksheets/sheet2.xml']=xml(doc)
 parts=dict(old.base.old.add_part(list(parts.items()),old.base.old.OPAQUE_NAME,old.base.old.OPAQUE,'application/octet-stream',('_rels/.rels',old.base.old.OPAQUE_NAME),'urn:smartdd:m9:opaque'))
 write(path,parts);return path

def format_candidate(root,parts):
 out=dict(parts);styles=tree(out['xl/styles.xml']);fonts=styles.find('{%s}fonts'%S);xfs=styles.find('{%s}cellXfs'%S)
 sheet=tree(out['xl/worksheets/sheet2.xml']);cells={x.get('r'):x for x in sheet.findall('.//{%s}c'%S)}
 source_style=int(cells['E2'].get('s','0'));xf=copy.deepcopy(xfs[source_style]);font=copy.deepcopy(fonts[int(xf.get('fontId','0'))])
 color=font.find('{%s}color'%S);color.set('rgb','FF008060');fonts.append(font);fonts.set('count',str(len(fonts)));xf.set('fontId',str(len(fonts)-1));xf.set('applyFont','1');xfs.append(xf);xfs.set('count',str(len(xfs)));cells['E2'].set('s',str(len(xfs)-1))
 # Merge is a display operation; both underlying cells survive in the bytes.
 merges=sheet.find('{%s}mergeCells'%S)
 if merges is None:merges=ET.SubElement(sheet,'{%s}mergeCells'%S,{'count':'1'})
 ET.SubElement(merges,'{%s}mergeCell'%S,{'ref':'G2:H2'})
 out['xl/styles.xml']=xml(styles);out['xl/worksheets/sheet2.xml']=xml(sheet)
 facts=guard(parts,out,{'xl/styles.xml','xl/worksheets/sheet2.xml'})
 path=root/'rich-formatted.xlsx';write(path,out);rw=load_workbook(path)
 assert rw['Calc']['E2'].value==123.456 and rw['Calc']['E2'].font.color.rgb=='FF008060'
 assert rw['Calc']['E3'].font.color.rgb=='FF123456' and rw['Calc']['E2'].number_format=='0.00'
 assert b'MERGE_COVERED' in out['xl/worksheets/sheet2.xml']
 facts['sharedStyleCopyOnWrite']=True;facts['coveredCellBytesPreserved']=True;return facts

def pptx_fixture(root):
 md=root/'rich-slides.md';md.write_text('# MASTER_TITLE\n\n| Label | Value |\n|---|---|\n| A | TABLE_OLD |\n| B | UNCHANGED_CELL |\n\n::: notes\nSLIDE_ONE_NOTES\n:::\n\n# OTHER_SLIDE\n\nUNCHANGED_SLIDE\n',encoding='utf8')
 path=root/'rich-source.pptx';subprocess.run(['pandoc',str(md),'--slide-level=1','-t','pptx','-o',str(path)],check=True,capture_output=True)
 parts=read(path)
 # Build an embedded spreadsheet + a real editable DrawingML chart.
 wb=Workbook();ws=wb.active;ws.title='Sheet1';ws.append(['Label','Value']);ws.append(['A',3]);ws.append(['B',4]);buf=BytesIO();wb.save(buf);embed=buf.getvalue()
 cx=f'''<c:chartSpace xmlns:c="{CH}" xmlns:a="{A}" xmlns:r="{O}"><c:chart><c:autoTitleDeleted val="1"/><c:plotArea><c:layout/><c:barChart><c:barDir val="col"/><c:grouping val="clustered"/><c:ser><c:idx val="0"/><c:order val="0"/><c:tx><c:v>Value</c:v></c:tx><c:cat><c:strRef><c:f>Sheet1!$A$2:$A$3</c:f><c:strCache><c:ptCount val="2"/><c:pt idx="0"><c:v>A</c:v></c:pt><c:pt idx="1"><c:v>B</c:v></c:pt></c:strCache></c:strRef></c:cat><c:val><c:numRef><c:f>Sheet1!$B$2:$B$3</c:f><c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="2"/><c:pt idx="0"><c:v>3</c:v></c:pt><c:pt idx="1"><c:v>4</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser><c:axId val="1"/><c:axId val="2"/></c:barChart><c:catAx><c:axId val="1"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="b"/><c:crossAx val="2"/><c:crosses val="autoZero"/></c:catAx><c:valAx><c:axId val="2"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:axPos val="l"/><c:crossAx val="1"/><c:crosses val="autoZero"/></c:valAx></c:plotArea><c:plotVisOnly val="1"/></c:chart><c:externalData r:id="rIdEmbedded"><c:autoUpdate val="0"/></c:externalData></c:chartSpace>'''.encode()
 parts['ppt/charts/chart1.xml']=cx;parts['ppt/embeddings/chart-data.xlsx']=embed
 parts['ppt/charts/_rels/chart1.xml.rels']=f'<Relationships xmlns="{R}"><Relationship Id="rIdEmbedded" Type="{O}/package" Target="../embeddings/chart-data.xlsx"/></Relationships>'.encode()
 slide=tree(parts['ppt/slides/slide1.xml']);sp=slide.find('.//{%s}spTree'%P);ids=[int(x.get('id')) for x in slide.findall('.//{%s}cNvPr'%P)];sid=max(ids)+1
 node=ET.fromstring(f'''<p:graphicFrame xmlns:p="{P}" xmlns:a="{A}" xmlns:c="{CH}" xmlns:r="{O}"><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="M9_CHART"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="5000000" y="2000000"/><a:ext cx="4000000" cy="2500000"/></p:xfrm><a:graphic><a:graphicData uri="{CH}"><c:chart r:id="rIdChart"/></a:graphicData></a:graphic></p:graphicFrame>''')
 sp.append(node);parts['ppt/slides/slide1.xml']=xml(slide)
 rel=tree(parts['ppt/slides/_rels/slide1.xml.rels']);ET.SubElement(rel,'{%s}Relationship'%R,{'Id':'rIdChart','Type':O+'/chart','Target':'../charts/chart1.xml'});parts['ppt/slides/_rels/slide1.xml.rels']=xml(rel)
 ct=tree(parts['[Content_Types].xml']);ET.SubElement(ct,'{%s}Override'%C,{'PartName':'/ppt/charts/chart1.xml','ContentType':'application/vnd.openxmlformats-officedocument.drawingml.chart+xml'});ET.SubElement(ct,'{%s}Override'%C,{'PartName':'/ppt/embeddings/chart-data.xlsx','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'});parts['[Content_Types].xml']=xml(ct)
 parts=dict(old.base.old.add_part(list(parts.items()),old.base.old.OPAQUE_NAME,old.base.old.OPAQUE,'application/octet-stream',('_rels/.rels',old.base.old.OPAQUE_NAME),'urn:smartdd:m9:opaque'))
 old.base.validate_relationships(parts);write(path,parts);return path

def pptx_patch(root,before):
 out=dict(before);slide=tree(out['ppt/slides/slide1.xml']);targets=[x for x in slide.findall('.//{%s}t'%A) if x.text=='TABLE_OLD'];assert len(targets)==1;targets[0].text='TABLE_EDITED'
 table=slide.find('.//{%s}tbl'%A);assert table is not None
 # Background local override; master/layout/theme remain byte-for-byte.
 cs=slide.find('{%s}cSld'%P);bg=ET.Element('{%s}bg'%P);props=ET.SubElement(bg,'{%s}bgPr'%P);fill=ET.SubElement(props,'{%s}solidFill'%A);ET.SubElement(fill,'{%s}srgbClr'%A,{'val':'F4FFF9'});cs.insert(0,bg)
 out['ppt/slides/slide1.xml']=xml(slide)
 chart=tree(out['ppt/charts/chart1.xml']);points=chart.findall('.//{%s}numCache/{%s}pt/{%s}v'%(CH,CH,CH));assert [x.text for x in points]==['3','4'];points[0].text='9';out['ppt/charts/chart1.xml']=xml(chart)
 with zipfile.ZipFile(BytesIO(out['ppt/embeddings/chart-data.xlsx'])) as z:emb={n:z.read(n) for n in z.namelist()}
 e=tree(emb['xl/worksheets/sheet1.xml']);cell=next(x for x in e.findall('.//{%s}c'%S) if x.get('r')=='B2');cell.find('{%s}v'%S).text='9';emb['xl/worksheets/sheet1.xml']=xml(e);buff=BytesIO()
 with zipfile.ZipFile(buff,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in emb.items():z.writestr(n,b)
 out['ppt/embeddings/chart-data.xlsx']=buff.getvalue()
 facts=guard(before,out,{'ppt/slides/slide1.xml','ppt/charts/chart1.xml','ppt/embeddings/chart-data.xlsx'})
 facts['mastersLayoutsThemesPreserved']={n:sha(b) for n,b in before.items() if n.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')) and b==out[n]}
 assert facts['mastersLayoutsThemesPreserved'];assert out['ppt/slides/slide2.xml']==before['ppt/slides/slide2.xml']
 write(root/'rich-edited.pptx',out);return facts

def main():
 p=argparse.ArgumentParser();p.add_argument('--prepare',required=True);args=p.parse_args();root=pathlib.Path(args.prepare).resolve();root.mkdir(parents=True,exist_ok=True)
 if (root/'preparation.json').exists():raise ValueError('use a fresh result directory; do not overwrite prior evidence')
 report={'createdAt':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'platform':platform.platform(),'scope':'generated Office samples and narrow adapter proofs; not product acceptance','sourceSha256':sha(HERE.read_bytes()),'checks':[]}
 def case(cid,title,fn):
  try:r=fn();report['checks'].append({'id':cid,'title':title,'status':'passed','facts':r})
  except Exception as e:report['checks'].append({'id':cid,'title':title,'status':'failed','error':repr(e)})
 # Exact prior failing fixture and exact prior deletion adapter, unchanged.
 src=old.fixture(root);parts=read(src);deleted,edits,removed=old.delete_sheet(parts,'Data');write(root/'oa12-deleted.xlsx',deleted)
 report['originalOa12']={'sourceSha256':sha(src.read_bytes()),'originalProbeSha256':sha(HERE.with_name('office-actions-probe.py').read_bytes()),'formulaEdits':edits,'removed':sorted(removed),'guard':guard(parts,deleted,{old.W,old.WR,'[Content_Types].xml','xl/worksheets/sheet2.xml'},removed=removed)}
 rich=rich_xlsx(root);case('OP01','XLSX styles copy-on-write, number semantics, merge retains source cell bytes',lambda:format_candidate(root,read(rich)))
 ppt=pptx_fixture(root);case('OP02','PPTX table + chart cache + embedded workbook edit; preserve masters/relationships',lambda:pptx_patch(root,read(ppt)))
 report['artifacts']={p.name:{'bytes':p.stat().st_size,'sha256':sha(p.read_bytes())} for p in root.iterdir() if p.is_file()}
 report['summary']={'passed':sum(x['status']=='passed' for x in report['checks']),'failed':sum(x['status']=='failed' for x in report['checks'])}
 (root/'preparation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report['summary']));raise SystemExit(bool(report['summary']['failed']))
if __name__=='__main__':main()
