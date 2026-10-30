"""Y10 real OOXML comments/revision anchors plus inert preservation payloads.
Uses existing installed fixture libraries only; no download or product changes.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib, json, tempfile, zipfile, xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

repo=Path(__file__).resolve().parents[2]
b=Path(tempfile.mkdtemp(prefix='willyou-office-preservation-9de5-',dir='/private/tmp'))
sha=lambda data:hashlib.sha256(data).hexdigest()
book=Workbook();sheet=book.active;sheet.title='Data'
sheet['A1']='KEEP_HEADER';sheet['A2']='OLD_XLSX_TEXT';sheet['B2']=3;sheet['B3']=4
sheet['A1'].comment=Comment('KEEP_XLSX_COMMENT','M9 Keep Author')
sheet['A2'].font=Font(name='Arial',size=14,bold=True,color='224466')
sheet['A2'].fill=PatternFill(fill_type='solid',fgColor='E0F0E8')
sheet.column_dimensions['A'].width=32
calc=book.create_sheet('Calc');calc['A1']='=SUM(Data!B2:B3)';calc['B1']='KEEP_OTHER_SHEET'
book.save(b/'seed.xlsx')
doc=Document();p=doc.add_paragraph('KEEP_DOCX_COMMENT_ANCHOR')
doc.add_comment(p.runs,text='KEEP_DOCX_COMMENT',author='M9 Keep Author',initials='M9')
p=doc.add_paragraph();ins=OxmlElement('w:ins')
for key,value in [('id','71'),('author','M9 Revision Author'),('date','2026-10-17T09:15:00Z')]:ins.set(qn('w:'+key),value)
run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='TRACKED_INSERT_KEEP';run.append(text);ins.append(run);p._p.append(ins)
doc.add_paragraph('DOCX_READONLY_BASELINE');doc.save(b/'seed.docx')
seed_pptx=repo/'docs/里程碑/M9_文件浏览工作区/验收/证据/predev-office/final/rich-source.pptx'
sources={'xlsx':b/'seed.xlsx','docx':b/'seed.docx','pptx':seed_pptx}
rows={}
for kind,path in sources.items():
    with zipfile.ZipFile(path)as z:parts={x.filename:z.read(x)for x in z.infolist()}
    if kind=='pptx':
        ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        slide=ET.fromstring(parts['ppt/slides/slide2.xml'])
        targets=[sp for sp in slide.findall('.//p:sp',ns)if 'OTHER_SLIDE' in ''.join(x.text or '' for x in sp.findall('.//a:t',ns))]
        assert len(targets)==1
        target={'slide':2,'name':targets[0].find('.//p:cNvPr',ns).attrib['name'],'part':'ppt/slides/slide2.xml','oldText':'OTHER_SLIDE','newText':'Y10_PPTX_EDIT'}
        parts['ppt/commentAuthors.xml']=b'<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cmAuthor id="0" name="M9 Keep Author" initials="M9" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>'
        parts['ppt/comments/comment1.xml']=b'<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cm authorId="0" dt="2026-10-17T09:15:00Z" idx="1"><p:pos x="10" y="20"/><p:text>KEEP_PPTX_COMMENT</p:text></p:cm></p:cmLst>'
        for part,relationship in [
            ('ppt/_rels/presentation.xml.rels',b'<Relationship Id="rIdY10Authors" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors" Target="commentAuthors.xml"/>'),
            ('ppt/slides/_rels/slide2.xml.rels',b'<Relationship Id="rIdY10Comments" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="../comments/comment1.xml"/>')]:
            assert b'</Relationships>'in parts[part];parts[part]=parts[part].replace(b'</Relationships>',relationship+b'</Relationships>')
        parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'</Types>',b'<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/><Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/></Types>')
    elif kind=='xlsx':target={'sheet':'Data','row':1,'column':0,'address':'A2','part':'xl/worksheets/sheet1.xml','oldText':'OLD_XLSX_TEXT','newText':'Y10_XLSX_EDIT'}
    else:target={'readOnly':True,'noApprovedManualEditing':True}
    extras={
      'customXml/y10-unknown.xml':b'<y10:unknown xmlns:y10="urn:willyou:m9:preservation" keep="true">KEEP_OPAQUE</y10:unknown>',
      'customXml/y10-revision-metadata.xml':b'<y10:metadata xmlns:y10="urn:willyou:m9:preservation" author="M9 Metadata Author" date="2026-10-17T09:15:00Z"/>',
      'customXml/y10-vbaProject.bin':b'INERT MACRO PLACEHOLDER ONLY; NO CODE OR EXECUTION\x00\xff'}
    for name in extras:assert name not in parts
    parts.update(extras)
    parts['_rels/.rels']=parts['_rels/.rels'].replace(b'</Relationships>',b'<Relationship Id="rIdY10Opaque" Type="urn:willyou:m9:preservation" Target="customXml/y10-unknown.xml"/></Relationships>')
    overrides=b''.join(f'<Override PartName="/{name}" ContentType="{("application/octet-stream"if name.endswith(".bin")else"application/xml")}"/>'.encode()for name in extras)
    parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'</Types>',overrides+b'</Types>')
    output=b/('preserve.'+kind)
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED)as z:
        for name,data in parts.items():z.writestr(name,data)
    with zipfile.ZipFile(output)as z:assert z.testzip()is None
    rows[kind]={'path':str(output),'sha256':sha(output.read_bytes()),'names':list(parts),'parts':{name:sha(data)for name,data in parts.items()},'target':target,'inertMacroPart':'customXml/y10-vbaProject.bin'}
result={'at':datetime.now(timezone.utc).isoformat(),'base':str(b),'files':rows,'fixtureDependencies':{'openpyxl':'existing3.1.5','python-docx':'existing1.2.0'},'primarySources':['https://learn.microsoft.com/en-us/office/open-xml/presentation/working-with-comments','https://learn.microsoft.com/en-us/office/open-xml/presentation/structure-of-a-presentationml-document','https://python-docx.readthedocs.io/en/latest/user/comments.html'],'boundary':'All generated public synthetic Office data. Real DOCX comment anchors/tracked insertion with author/date, real XLSX comments/VML and PPTX comment author/slide relationship metadata. Custom XML revision metadata is opaque preservation only, not native spreadsheet/slide tracked revision support. Macro payload is explicitly inert placeholder, not a real executable macro or macro security certification. DOCX readonly verification is authorized separately; original Y10 manual-DOCX edit contract remains pending, not silently replaced.'}
(b/'fixtures.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'manifest':str(b/'fixtures.json'),'base':str(b),'files':list(rows)}))
