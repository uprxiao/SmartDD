"""Independent ZIP/XML + existing openpyxl/python-docx readers, not App models."""
from pathlib import Path
import copy, datetime, hashlib, io, json, sys, zipfile, xml.etree.ElementTree as ET
from openpyxl import load_workbook
from docx import Document

b=Path(sys.argv[1]);r=json.loads((b/'result.json').read_text())
assert r['status']=='selected-normal-text-edits-and-approved-DOCX-readonly-passed;formal-Y10-pending'
assert len(r['checks'])==6
sha=lambda data:hashlib.sha256(data).hexdigest()
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def package(p):
    with zipfile.ZipFile(p)as z:
        assert z.testzip()is None
        return z.namelist(),{n:z.read(n)for n in z.namelist()}
def canonical(node):
    node=copy.deepcopy(node)
    for x in node.iter():
        attrs=dict(sorted(x.attrib.items()));x.attrib.clear();x.attrib.update(attrs)
    return ET.tostring(node,encoding='unicode')
checks=[]
for c in r['checks']:
    name=c['theme']+'.'+c['kind'];source=r['corpus'][name];path=Path(source['path']);fixture=source['fixture']
    names,before=package(fixture['path']);afterNames,after=package(path)
    assert names==afterNames==fixture['names'],name
    changed=[n for n in names if before[n]!=after[n]]
    assert sha(after[fixture['inertMacroPart']])==fixture['parts'][fixture['inertMacroPart']]
    if c['kind']=='docx':
        stat=path.stat();assert {'sha256':sha(path.read_bytes()),'inode':str(stat.st_ino),'mtimeNs':str(stat.st_mtime_ns)}==source['before'];assert changed==[]
        doc=Document(path);assert len(doc.comments)==1
        comment=list(doc.comments)[0];assert comment.author=='M9 Keep Author'and comment.text=='KEEP_DOCX_COMMENT'
        main=ET.fromstring(after['word/document.xml']);revisions=main.findall('.//w:ins',ns);assert len(revisions)==1
        assert revisions[0].attrib['{'+ns['w']+'}author']=='M9 Revision Author'
        assert revisions[0].attrib['{'+ns['w']+'}date']=='2026-10-17T09:15:00Z'
        ids=lambda tag:[x.attrib['{'+ns['w']+'}id']for x in main.findall('.//w:'+tag,ns)]
        assert ids('commentRangeStart')==ids('commentRangeEnd')==ids('commentReference')==['0']
        checks.append({'file':name,'branch':'approved-readonly-not-original-Y10-edit','allPartsAndSourceStatExact':True,'commentAnchorsAuthorAndTrackedRevisionDateVerified':True});continue
    target=fixture['target']
    if c['kind']=='xlsx':
        allowed={'xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml','xl/workbook.xml'}
        assert set(changed)<=allowed and target['part']in changed
        for part in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
            old=ET.fromstring(before[part]);new=ET.fromstring(after[part]);a={x.attrib['r']:x for x in old.findall('.//s:c',ns)};z={x.attrib['r']:x for x in new.findall('.//s:c',ns)};assert a.keys()==z.keys()
            for address,cell in a.items():
                current=z[address]
                if part==target['part']and address==target['address']:
                    assert current.attrib.get('s')==cell.attrib.get('s')and current.attrib['r']==cell.attrib['r']
                    assert ''.join(x.text or ''for x in current.findall('.//s:t',ns))==target['newText']
                    continue
                left=copy.deepcopy(cell);right=copy.deepcopy(current)
                if cell.find('s:f',ns)is not None:
                    assert canonical(cell.find('s:f',ns))==canonical(current.find('s:f',ns))
                    for x in [left,right]:
                        x.attrib.pop('t',None)
                        for child in list(x.findall('s:v',ns)):x.remove(child)
                    assert part=='xl/worksheets/sheet2.xml'and address=='A1'
                    assert float(current.find('s:v',ns).text)==7
                assert canonical(left)==canonical(right),(name,part,address)
            for tree in [old,new]:
                for parent in tree.iter():
                    for child in list(parent):
                        if child.tag=='{'+ns['s']+'}c':parent.remove(child)
            assert canonical(old)==canonical(new),(name,part,'noncell')
        old=ET.fromstring(before['xl/workbook.xml']);new=ET.fromstring(after['xl/workbook.xml'])
        for tree in [old,new]:
            for x in list(tree.findall('s:calcPr',ns)):tree.remove(x)
        assert canonical(old)==canonical(new)
        reader=load_workbook(path,data_only=False);assert reader['Data']['A2'].value==target['newText']
        assert reader['Data']['A1'].comment.text=='KEEP_XLSX_COMMENT'and reader['Data']['A1'].comment.author=='M9 Keep Author'
        assert load_workbook(path,data_only=True)['Calc']['A1'].value==7
    else:
        assert changed==[target['part']]
        old=ET.fromstring(before[target['part']]);new=ET.fromstring(after[target['part']])
        texts=[x for x in new.findall('.//a:t',ns)if x.text==target['newText']];assert len(texts)==1
        texts[0].text=target['oldText'];assert canonical(old)==canonical(new)
        comment=ET.fromstring(after['ppt/comments/comment1.xml']).find('p:cm',ns)
        assert comment.attrib=={'authorId':'0','dt':'2026-10-17T09:15:00Z','idx':'1'}
        assert comment.find('p:text',ns).text=='KEEP_PPTX_COMMENT'
    checks.append({'file':name,'branch':'actual-allowed-visible-text-edit','changedParts':changed,'namesAndOrderExact':True,'allUneditedPartsExact':True,'editedNodeFormattingAndOtherSemanticsPreserved':True,'realCommentRelationshipsAndOpaquePartsExact':True,'sourceHash':sha(path.read_bytes())})
out={'at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'status':'passed-selected-preservation-readers;formal-Y10-pending','source':str(b/'result.json'),'checks':checks,'boundary':'Independent Python ZIP CRC/ElementTree plus existing openpyxl/python-docx readers.4 actual allowed UI text edits and2 approved DOCX readonly checks only. Valid formula cache changes explicitly constrained, no any-rewrite rejection. No original manual-DOCX edit, LibreOffice-render, executable-macro security, other-platform or whole-Y10 completion claim.'}
(b/'independent-preservation.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'status':out['status'],'files':len(checks),'edited':4,'readonlyDOCX':2}))
